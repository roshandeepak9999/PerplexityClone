"""
Perplexity Clone - Flask Backend
---------------------------------
Modules:
  1. Auth module        - register / login / logout (Flask-Login + SQLite)
  2. Agent module        - plans which tool fits a plain-text question:
                            search / calculate / direct (agent.py)
  3. Search module       - DuckDuckGo web search (with retries)
  4. RAG module          - page fetch, chunking, embeddings, similarity
                            search over chunks (rag.py) - real classical RAG
  5. LLM answer module    - Gemini, grounded in retrieved chunks / search
                            results, with a general-knowledge fallback
  6. Guardrail module    - a post-hoc pass checks whether the answer's
                            claims are backed by the retrieved sources
  7. History module       - every search a logged-in user makes is saved
  8. Related questions module

Author: <your name>
Project: MCA Short-Term Internship - Generative AI Engineering
"""

import io
import os
import re
import time
import uuid

from dotenv import load_dotenv
from werkzeug.exceptions import HTTPException
from flask import Flask, render_template, request, jsonify, redirect, url_for, flash
from flask_login import (
    LoginManager,
    login_user,
    logout_user,
    login_required,
    current_user,
)

# The package was renamed: `duckduckgo_search` -> `ddgs`.
# Prefer the new one, fall back to the old one if that's what is installed.
try:
    from ddgs import DDGS
except ImportError:  # pragma: no cover
    from duckduckgo_search import DDGS

from google import genai
from google.genai import types

from models import db, User, SearchHistory
from agent import plan_tool, try_calculate
from rag import build_chunk_store, build_context_from_chunks

# ---------------------------------------------------------------------------
# Setup
# ---------------------------------------------------------------------------
load_dotenv()

app = Flask(__name__)
app.secret_key = os.getenv("FLASK_SECRET_KEY", "dev-secret-key-change-this")
app.config["SQLALCHEMY_DATABASE_URI"] = "sqlite:///perplexed.db"
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False
app.config["MAX_CONTENT_LENGTH"] = 20 * 1024 * 1024  # 20 MB per request

# Profile photos are stored on disk here and served by Flask's static route.
AVATAR_DIR = os.path.join(app.static_folder, "uploads", "avatars")
os.makedirs(AVATAR_DIR, exist_ok=True)
MAX_AVATAR_BYTES = 5 * 1024 * 1024  # 5 MB

db.init_app(app)

login_manager = LoginManager()
login_manager.init_app(app)
login_manager.login_view = "login"


@login_manager.user_loader
def load_user(user_id):
    return db.session.get(User, int(user_id))


GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")
MODEL_NAME = "gemini-flash-latest"

client = genai.Client(api_key=GEMINI_API_KEY) if GEMINI_API_KEY else None


def ensure_user_columns():
    """create_all() never alters existing tables, so add the new profile
    columns to an older database automatically (keeps your data)."""
    from sqlalchemy import text

    with db.engine.begin() as conn:
        cols = {row[1] for row in conn.execute(text("PRAGMA table_info(users)"))}
        if "full_name" not in cols:
            conn.execute(text("ALTER TABLE users ADD COLUMN full_name VARCHAR(120)"))
        if "photo" not in cols:
            conn.execute(text("ALTER TABLE users ADD COLUMN photo VARCHAR(255)"))


with app.app_context():
    db.create_all()
    ensure_user_columns()


# ---------------------------------------------------------------------------
# Core logic
# ---------------------------------------------------------------------------
def web_search(query, max_results=6, retries=3):
    """Search the web using DuckDuckGo. Retries a few times because DDG
    sometimes rate-limits or returns an empty list on the first attempt."""
    for attempt in range(1, retries + 1):
        try:
            with DDGS() as ddgs:
                raw = list(ddgs.text(query, max_results=max_results))

            results = [
                {
                    "title": r.get("title", "Untitled"),
                    "url": r.get("href", ""),
                    "snippet": r.get("body", ""),
                }
                for r in raw
                if r.get("href")
            ]
            print(f"[web_search] attempt {attempt}: {len(results)} results for {query!r}")
            if results:
                return results
        except Exception as exc:
            print(f"[web_search] attempt {attempt} error: {exc}")

        time.sleep(1.5 * attempt)  # simple backoff

    return []


def build_context(results):
    blocks = []
    for i, r in enumerate(results, start=1):
        blocks.append(f"[{i}] {r['title']}\n{r['snippet']}\nSource: {r['url']}")
    return "\n\n".join(blocks)


FORMAT_GUIDE = """Choose whichever answer format best fits the QUESTION TYPE, don't force one
style on everything:
- "How to" / process / step-by-step questions -> a numbered list of steps.
- Comparisons ("X vs Y", pros and cons) -> a short markdown table or bullet
  points grouped per item.
- "What is" / conceptual / explain questions -> a short paragraph, optionally
  followed by bullet points for key facts.
- Lists ("top 5...", "examples of...") -> a numbered or bulleted list.
- Yes/no or factual lookups -> a direct one- or two-sentence answer first,
  then brief supporting detail."""


def is_failed_answer(text):
    """True for error / 'no info' answers we don't want to reuse as context."""
    if not text:
        return True
    t = text.lower()
    return (
        t.startswith("error generating answer")
        or t.startswith("⚠️")
        or "do not contain any information" in t
        or "do not contain enough information" in t
    )



# ---------------------------------------------------------------------------
# File / image uploads
# ---------------------------------------------------------------------------
MAX_FILES = 5
MAX_TEXT_CHARS = 20000  # per text/docx file, to keep prompts small

IMAGE_TYPES = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".webp": "image/webp",
    ".gif": "image/gif",
}
TEXT_EXTS = {
    ".txt", ".md", ".csv", ".json", ".xml", ".yaml", ".yml", ".log",
    ".py", ".js", ".html", ".css", ".java", ".c", ".cpp", ".sql",
}


def process_uploads(files):
    """Turn uploaded files into something Gemini can read.

    Returns None when there are no files, otherwise a dict:
      parts: inline binary parts (images, PDFs) sent straight to Gemini
      texts: [(filename, extracted_text)] for text-like files and .docx
      names: original filenames
    Raises ValueError with a user-friendly message on bad input.
    Files are processed in memory only; nothing is saved to disk.
    """
    if not files:
        return None
    if len(files) > MAX_FILES:
        raise ValueError(f"You can attach up to {MAX_FILES} files at a time.")

    parts, texts, names = [], [], []
    for f in files:
        name = os.path.basename(f.filename)
        ext = os.path.splitext(name)[1].lower()
        data = f.read()
        if not data:
            raise ValueError(f"'{name}' is empty.")

        if ext in IMAGE_TYPES:
            parts.append({"mime_type": IMAGE_TYPES[ext], "data": data})
        elif ext == ".pdf":
            parts.append({"mime_type": "application/pdf", "data": data})
        elif ext == ".docx":
            try:
                from docx import Document
            except ImportError:
                raise ValueError("Word files need: pip install python-docx")
            doc = Document(io.BytesIO(data))
            text = "\n".join(p.text for p in doc.paragraphs)
            texts.append((name, text[:MAX_TEXT_CHARS]))
        elif ext in TEXT_EXTS:
            text = data.decode("utf-8", errors="replace")
            texts.append((name, text[:MAX_TEXT_CHARS]))
        else:
            raise ValueError(
                f"'{name}' is not a supported file type. "
                "Use images, PDF, DOCX, or text/code files."
            )
        names.append(name)

    return {"parts": parts, "texts": texts, "names": names}


# ---------------------------------------------------------------------------
# Coding mode
# ---------------------------------------------------------------------------
_LANGS = (
    r"python|java|javascript|typescript|c\+\+|c#|golang|rust|php|ruby|kotlin|swift|"
    r"html|css|sql|react|node|flask|django|bash|powershell|dart|flutter"
)
_CODE_RE = re.compile(
    r"\b(code|coding|program|script|function|snippet|algorithm|regex|query|"
    r"implement|debug|refactor|"
    r"(develop|build|create|make|write|generate)\b.{0,40}\b(app|website|web ?page|api|bot|"
    r"clone|game|class|component|tool|system|project|server|login|page)|"
    rf"in ({_LANGS}))\b",
    re.IGNORECASE,
)


def is_code_request(query):
    """Heuristic: does the user want code written (vs. a factual answer)?"""
    return bool(_CODE_RE.search(query))


CODE_GUIDE = """You are an expert software engineer. The user wants CODE.
- Give complete, working, runnable code. Never use placeholders such as
  "...", "rest of the code here" or "add your logic here".
- Put every file in its own fenced markdown code block with the language tag
  (```python, ```html, ...), and write the file name on the line above it.
- Start with 1-2 sentences on the approach. If there are several files, show
  the project structure first.
- After the code, add short "How to run" steps and any packages to install.
- Use the language the user asked for; if none is given, pick the most suitable
  one and say which.
- Do not add [n] citations. Keep explanations brief; the code is the answer."""


# Models to try, in order. If the first is overloaded (503), the lighter
# fallback usually still responds.
MODELS = [MODEL_NAME, "gemini-flash-lite-latest"]
TRANSIENT_MARKERS = ("503", "429", "UNAVAILABLE", "RESOURCE_EXHAUSTED", "overloaded")


def call_gemini(contents, retries=3):
    """Call Gemini, retrying temporary failures (503 overloaded / 429 rate
    limit) with backoff, then falling back to the next model in MODELS."""
    last_exc = None
    for model in MODELS:
        for attempt in range(retries):
            try:
                response = client.models.generate_content(
                    model=model,
                    contents=contents,
                    config=types.GenerateContentConfig(max_output_tokens=16384),
                )
                return (response.text or "").strip()
            except Exception as exc:
                last_exc = exc
                transient = any(m in str(exc) for m in TRANSIENT_MARKERS)
                print(f"[gemini] {model} attempt {attempt + 1} failed: {str(exc)[:120]}")
                if not transient:
                    raise
                if attempt < retries - 1:
                    time.sleep(2 ** attempt)  # 1s, 2s
    raise last_exc


def generate_answer(query, results, history=None, attachments=None, code_mode=False,
                     context_override=None):
    """Answer using the search results (or a RAG context_override) when
    available; otherwise fall back to the model's own knowledge (and say so).

    context_override: a pre-built, numbered context string (e.g. from the
    RAG chunk/embedding pipeline in rag.py) to use INSTEAD of raw snippet
    context built from `results`. Citations [n] still refer to the same
    source numbers either way.
    """
    if not GEMINI_API_KEY:
        return (
            "⚠️ No GEMINI_API_KEY configured. Add one to your .env file "
            "(see README.md) to enable AI-generated answers."
        )

    history_text = ""
    if history:
        for h in history[-4:]:
            history_text += f"Previous Q: {h['query']}\nPrevious A: {h['answer']}\n\n"

    if attachments:
        attach_text = "\n\n".join(
            f"--- File: {n} ---\n{t}" for n, t in attachments["texts"]
        )
        prompt = f"""You are a helpful AI assistant, similar to Perplexity AI.
The user attached one or more files (documents and/or images). Answer their
message using the attached files as the primary source. If the files don't
contain what is needed, say so briefly and then answer from general knowledge
if you reasonably can.

{FORMAT_GUIDE}

Be concise and factual.

{history_text}{attach_text}

User message: {query}

Answer:"""
    elif code_mode:
        prompt = f"""{CODE_GUIDE}

{history_text}User request: {query}

Answer:"""
    elif results or context_override:
        context = context_override if context_override else build_context(results)
        prompt = f"""You are an AI search assistant, similar to Perplexity AI.
Answer the user's question primarily using the numbered search results below.
Cite the source(s) you used inline with their number, e.g. [1], directly after
the relevant sentence.

If the results only partly cover the question, answer the parts they support
with citations, then fill any gap from your general knowledge WITHOUT a citation
and say briefly that it comes from general knowledge. Never reply that there is
no information if you can reasonably answer the question.

{FORMAT_GUIDE}

Be concise and factual.

{history_text}Search results:
{context}

Question: {query}

Answer (with inline [n] citations where sources are used):"""
    else:
        prompt = f"""You are an AI assistant, similar to Perplexity AI. Live web search
returned no results for this question, so answer from your own general knowledge.
Start with one short line saying no live sources were found, then answer the
question clearly. Do not include [n] citations.

{FORMAT_GUIDE}

Be concise and factual.

{history_text}Question: {query}

Answer:"""

    try:
        contents = [prompt]
        if attachments:
            contents += [
                types.Part.from_bytes(data=p["data"], mime_type=p["mime_type"])
                for p in attachments["parts"]
            ]
        return call_gemini(contents)
    except Exception as exc:
        if any(m in str(exc) for m in TRANSIENT_MARKERS):
            return (
                "Error generating answer: the AI service is busy right now. "
                "Please try again in a minute."
            )
        return f"Error generating answer: {exc}"


def generate_direct_answer(query, history=None):
    """For questions the agent decided don't need web retrieval at all -
    a clean, confident answer with no 'no live sources found' framing,
    since search was never attempted."""
    if not GEMINI_API_KEY:
        return (
            "⚠️ No GEMINI_API_KEY configured. Add one to your .env file "
            "(see README.md) to enable AI-generated answers."
        )

    history_text = ""
    if history:
        for h in history[-4:]:
            history_text += f"Previous Q: {h['query']}\nPrevious A: {h['answer']}\n\n"

    prompt = f"""You are a helpful AI assistant. Answer the question directly
from your own knowledge - this question doesn't need a live web search.

{FORMAT_GUIDE}

{history_text}Question: {query}

Answer:"""

    try:
        return call_gemini(prompt)
    except Exception as exc:
        if any(m in str(exc) for m in TRANSIENT_MARKERS):
            return (
                "Error generating answer: the AI service is busy right now. "
                "Please try again in a minute."
            )
        return f"Error generating answer: {exc}"


def verify_answer(answer, context):
    """Hallucination guardrail: a SEPARATE model call that checks whether
    the answer's claims are actually backed by the retrieved context.
    Returns 'supported', 'unsupported', or None (skipped/unavailable)."""
    if not GEMINI_API_KEY or not context or is_failed_answer(answer):
        return None

    prompt = f"""You are a strict fact-checker. Compare the ANSWER to the
SOURCES below. Reply with EXACTLY ONE WORD:
"SUPPORTED" if every factual claim in the ANSWER is backed by the SOURCES,
or "UNSUPPORTED" if the ANSWER includes claims not found in the SOURCES.

SOURCES:
{context}

ANSWER:
{answer}

Reply with one word only:"""

    try:
        verdict = call_gemini(prompt, retries=1).strip().upper()
        if "UNSUPPORTED" in verdict:
            return "unsupported"
        if "SUPPORTED" in verdict:
            return "supported"
        return None
    except Exception:
        return None


def generate_related_questions(query, answer):
    if not GEMINI_API_KEY or is_failed_answer(answer):
        return []

    prompt = f"""Based on the question and answer below, suggest exactly 3 short,
relevant follow-up questions a curious user might ask next. Return ONLY the
3 questions, one per line, with no numbering or extra text.

Question: {query}
Answer: {answer}"""

    try:
        text = call_gemini(prompt, retries=2)
        lines = [
            line.strip("-•* ").strip()
            for line in text.split("\n")
            if line.strip()
        ]
        return lines[:3]
    except Exception:
        return []


def get_session_history(limit=10):
    """Recent *successful* history for the logged-in user, used as LLM context.
    Failed / 'no info' answers are skipped so they don't poison new answers."""
    if not current_user.is_authenticated:
        return []
    rows = (
        SearchHistory.query.filter_by(user_id=current_user.id)
        .order_by(SearchHistory.created_at.desc())
        .limit(limit)
        .all()
    )
    return [
        {"query": r.query_text, "answer": r.answer}
        for r in reversed(rows)
        if not is_failed_answer(r.answer)
    ]


# ---------------------------------------------------------------------------
# Auth routes
# ---------------------------------------------------------------------------
@app.route("/register", methods=["GET", "POST"])
def register():
    if current_user.is_authenticated:
        return redirect(url_for("index"))

    if request.method == "POST":
        username = request.form.get("username", "").strip()
        email = request.form.get("email", "").strip().lower()
        password = request.form.get("password", "")

        if not username or not email or not password:
            flash("Please fill in all fields.", "error")
            return render_template("register.html")

        if User.query.filter(
            (User.username == username) | (User.email == email)
        ).first():
            flash("Username or email already registered.", "error")
            return render_template("register.html")

        user = User(username=username, email=email)
        user.set_password(password)
        db.session.add(user)
        db.session.commit()

        login_user(user)
        return redirect(url_for("index"))

    return render_template("register.html")


@app.route("/login", methods=["GET", "POST"])
def login():
    if current_user.is_authenticated:
        return redirect(url_for("index"))

    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")

        user = User.query.filter_by(username=username).first()
        if user and user.check_password(password):
            login_user(user)
            return redirect(url_for("index"))

        flash("Invalid username or password.", "error")
        return render_template("login.html")

    return render_template("login.html")


@app.route("/logout")
@login_required
def logout():
    logout_user()
    return redirect(url_for("login"))


# ---------------------------------------------------------------------------
# Main routes
# ---------------------------------------------------------------------------
@app.route("/")
@login_required
def index():
    return render_template("index.html", user=current_user)


@app.route("/api/search", methods=["POST"])
@login_required
def api_search():
    # Accept multipart/form-data (with files) or plain JSON (text only)
    if request.files or request.form:
        query = (request.form.get("query") or "").strip()
        files = [f for f in request.files.getlist("files") if f and f.filename]
    else:
        data = request.get_json(silent=True) or {}
        query = (data.get("query") or "").strip()
        files = []

    if not query and not files:
        return jsonify({"error": "Empty query"}), 400
    if not query:
        query = "Describe and summarize the attached file(s)."

    try:
        attachments = process_uploads(files)
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400

    history = get_session_history()

    # Coding requests and attachments don't need web search - unchanged
    # from before, and they SKIP the agent/RAG pipeline entirely.
    code_mode = is_code_request(query) and not attachments

    results = []
    tool_used = None
    verified = None

    if attachments:
        tool_used = "attachment"
        answer = generate_answer(query, results, history, attachments, code_mode)
    elif code_mode:
        tool_used = "code"
        answer = generate_answer(query, results, history, attachments, code_mode)
    else:
        # ---- Agent: PLAN which tool fits this plain-text question ----
        tool_used = plan_tool(call_gemini, query) if GEMINI_API_KEY else "search"

        if tool_used == "calculate":
            calc_result = try_calculate(query)
            if calc_result is not None:
                answer = f"**{calc_result}**"
            else:
                tool_used = "search"  # couldn't parse as math -> fall back

        if tool_used == "direct":
            answer = generate_direct_answer(query, history)

        if tool_used == "search":
            results = web_search(query)

            # ---- RAG: chunk + embed + similarity search over results ----
            chunks = build_chunk_store(client, query, results)
            context_override = build_context_from_chunks(chunks) if chunks else None

            answer = generate_answer(
                query, results, history, context_override=context_override
            )

            # ---- Guardrail: verify the answer against retrieved context ----
            check_context = context_override or (build_context(results) if results else None)
            verified = verify_answer(answer, check_context)

    related = generate_related_questions(query, answer)

    saved_query = query
    if attachments:
        saved_query = f"{query}  \U0001F4CE {', '.join(attachments['names'])}"[:500]

    entry = SearchHistory(
        user_id=current_user.id,
        query=saved_query,  # mapped to query_text by the model
        answer=answer,
        sources=results,    # stored as JSON by the model
    )
    db.session.add(entry)
    db.session.commit()

    return jsonify(
        {
            "query": query,
            "answer": answer,
            "sources": results,
            "related_questions": related,
            "history_id": entry.id,
            "attachments": attachments["names"] if attachments else [],
            "tool_used": tool_used,
            "verified": verified,
        }
    )


@app.route("/api/history", methods=["GET"])
@login_required
def api_history():
    rows = (
        SearchHistory.query.filter_by(user_id=current_user.id)
        .order_by(SearchHistory.created_at.desc())
        .limit(50)
        .all()
    )
    return jsonify([r.to_dict() for r in rows])


@app.route("/api/history/<int:history_id>", methods=["DELETE"])
@login_required
def delete_history_item(history_id):
    entry = SearchHistory.query.filter_by(
        id=history_id, user_id=current_user.id
    ).first()
    if not entry:
        return jsonify({"error": "Not found"}), 404
    db.session.delete(entry)
    db.session.commit()
    return jsonify({"status": "deleted"})


@app.route("/api/new-chat", methods=["POST"])
@login_required
def new_chat():
    # "New chat" just starts a fresh conversation thread in the UI;
    # past searches remain saved in history.
    return jsonify({"status": "ok"})


# ---------------------------------------------------------------------------
# Profile: personal details, photo, password
# ---------------------------------------------------------------------------
def sniff_image_ext(data):
    """Detect the real image type from the file's first bytes (ignores the
    filename), so only genuine PNG/JPG/GIF/WEBP files are accepted."""
    if data.startswith(b"\x89PNG\r\n\x1a\n"):
        return "png"
    if data[:3] == b"\xff\xd8\xff":
        return "jpg"
    if data[:6] in (b"GIF87a", b"GIF89a"):
        return "gif"
    if data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return "webp"
    return None


def delete_avatar_file(filename):
    if not filename:
        return
    try:
        os.remove(os.path.join(AVATAR_DIR, os.path.basename(filename)))
    except OSError:
        pass


def save_avatar(user, file_storage):
    data = file_storage.read()
    if not data:
        raise ValueError("That file is empty.")
    if len(data) > MAX_AVATAR_BYTES:
        raise ValueError("Photo is too large (5 MB max).")
    ext = sniff_image_ext(data)
    if not ext:
        raise ValueError("Please upload a PNG, JPG, WEBP or GIF image.")

    filename = f"u{user.id}_{uuid.uuid4().hex[:10]}.{ext}"
    with open(os.path.join(AVATAR_DIR, filename), "wb") as fh:
        fh.write(data)

    delete_avatar_file(user.photo)  # remove the previous photo
    user.photo = filename


@app.route("/profile")
@login_required
def profile():
    return render_template("profile.html", user=current_user)


@app.route("/profile/details", methods=["POST"])
@login_required
def profile_details():
    full_name = request.form.get("full_name", "").strip()[:120]
    username = request.form.get("username", "").strip()
    email = request.form.get("email", "").strip().lower()

    if not username or not email:
        flash("Username and email are required.", "error")
        return redirect(url_for("profile"))
    if len(username) > 80 or len(email) > 120:
        flash("Username or email is too long.", "error")
        return redirect(url_for("profile"))
    if "@" not in email or "." not in email.split("@")[-1]:
        flash("Please enter a valid email address.", "error")
        return redirect(url_for("profile"))

    clash = User.query.filter(
        ((User.username == username) | (User.email == email))
        & (User.id != current_user.id)
    ).first()
    if clash:
        flash("That username or email is already taken.", "error")
        return redirect(url_for("profile"))

    current_user.full_name = full_name or None
    current_user.username = username
    current_user.email = email
    db.session.commit()
    flash("Personal details updated.", "success")
    return redirect(url_for("profile"))


@app.route("/profile/photo", methods=["POST"])
@login_required
def profile_photo():
    f = request.files.get("photo")
    if not f or not f.filename:
        flash("Choose a photo first.", "error")
        return redirect(url_for("profile"))
    try:
        save_avatar(current_user, f)
        db.session.commit()
        flash("Profile photo updated.", "success")
    except ValueError as exc:
        flash(str(exc), "error")
    return redirect(url_for("profile"))


@app.route("/profile/photo/remove", methods=["POST"])
@login_required
def profile_photo_remove():
    delete_avatar_file(current_user.photo)
    current_user.photo = None
    db.session.commit()
    flash("Profile photo removed.", "success")
    return redirect(url_for("profile"))


@app.route("/profile/password", methods=["POST"])
@login_required
def profile_password():
    current = request.form.get("current_password", "")
    new = request.form.get("new_password", "")
    confirm = request.form.get("confirm_password", "")

    if not current_user.check_password(current):
        flash("Current password is incorrect.", "error")
    elif len(new) < 8:
        flash("New password must be at least 8 characters.", "error")
    elif new != confirm:
        flash("New passwords don't match.", "error")
    else:
        current_user.set_password(new)
        db.session.commit()
        flash("Password updated.", "success")
    return redirect(url_for("profile"))


@app.errorhandler(Exception)
def handle_error(e):
    """Return JSON for API errors so the frontend shows the real message
    instead of failing with "Unexpected token '<'"."""
    code = e.code if isinstance(e, HTTPException) else 500
    if not isinstance(e, HTTPException):
        app.logger.exception(e)

    if request.path.startswith("/api/"):
        if code == 413:
            return jsonify({"error": "Files are too large (20 MB max in total)."}), 413
        return jsonify({"error": str(e)}), code

    if isinstance(e, HTTPException):
        return e  # normal 404/403 pages
    raise e  # real crashes: Flask's debug page


if __name__ == "__main__":
    app.run(debug=True, port=5000)